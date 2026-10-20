import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

SAMPLE_DIR = os.path.dirname(__file__)

def generate_excel_sample(filepath: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Al Rajhi Statement"

    # Top metadata
    ws.cell(row=1, column=1, value="Al Rajhi Bank - Corporate Account Statement").font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    ws.cell(row=2, column=1, value="Company Name: Horizon Global Trading Ltd.").font = Font(name="Segoe UI", size=10, bold=True)
    ws.cell(row=3, column=1, value="Account Number: 1002345678901 | IBAN: SA03800001002345678901").font = Font(name="Segoe UI", size=10)
    ws.cell(row=4, column=1, value="Statement Period: 01/10/2026 - 20/10/2026 | Currency: SAR").font = Font(name="Segoe UI", size=10, italic=True)

    headers = ["Date", "Description (البيان)", "Reference", "Debit (مدين)", "Credit (دائن)", "Balance (الرصيد)"]
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows = [
        ("2026-10-01", "Opening Balance", "REF-001", 0.0, 0.0, 500000.0),
        ("2026-10-02", "Interbank Transfer to SNB Treasury Acct 2009876543210", "REF-002", 75000.0, 0.0, 425000.0),
        ("2026-10-03", "تحويل رواتب - محمد عبدالله العتيبي Emp 1001", "REF-003", 12500.0, 0.0, 412500.0),
        ("2026-10-04", "سداد فاتورة 014 الجوازات وتأشيرات MOI Absher", "REF-004", 4200.0, 0.0, 408300.0),
        ("2026-10-05", "سداد دفعة عقد إيجار منصة إيجار EJ-987456", "REF-005", 45000.0, 0.0, 363300.0),
        ("2026-10-06", "Customer Payment - Al Safa Trading", "REF-006", 0.0, 120000.0, 483300.0),
        ("2026-10-07", "Inter-Bank Inward Transfer from Riyad Bank 3001122334455", "REF-007", 0.0, 50000.0, 533300.0),
        ("2026-10-08", "Office Supplies & Stationery", "REF-008", 1850.0, 0.0, 531450.0),
        ("2026-10-09", "تحويل راتب الموظفة سارة فهد الشهري Emp 1004", "REF-009", 9800.0, 0.0, 521650.0),
        ("2026-10-10", "SADAD Payment 044 ZATCA VAT Return Payment", "REF-010", 15400.0, 0.0, 506250.0)
    ]

    for r_idx, r_data in enumerate(rows, 7):
        for c_idx, val in enumerate(r_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if c_idx in [4, 5, 6]:
                cell.number_format = '#,##0.00'

    for col in ws.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 24

    wb.save(filepath)

def generate_xml_sample(filepath: str):
    xml_content = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.02">
  <BkToCstmrStmt>
    <Stmt>
      <Id>SNB-STMT-2026-10</Id>
      <Acct>
        <Id>
          <IBAN>SA441000002009876543210</IBAN>
        </Id>
      </Acct>
      <!-- Entry 1: Inter-bank credit from Al Rajhi -->
      <Ntry>
        <Amt Ccy="SAR">75000.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <BookgDt><Dt>2026-10-02</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs><EndToEndId>TXN-SNB-101</EndToEndId></Refs>
            <RmtInf>
              <Ustrd>Interbank Inward Transfer from Al Rajhi Bank 1002345678901 Treasury</Ustrd>
            </RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
      <!-- Entry 2: Payroll to Ahmed Ali Al-Ghamdi -->
      <Ntry>
        <Amt Ccy="SAR">11000.00</Amt>
        <CdtDbtInd>DBIT</CdtDbtInd>
        <BookgDt><Dt>2026-10-03</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs><EndToEndId>TXN-SNB-102</EndToEndId></Refs>
            <RmtInf>
              <Ustrd>Payroll Salary Transfer - Ahmed Ali Al-Ghamdi Emp 1002</Ustrd>
            </RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
      <!-- Entry 3: SADAD Electricity SEC 002 -->
      <Ntry>
        <Amt Ccy="SAR">6850.00</Amt>
        <CdtDbtInd>DBIT</CdtDbtInd>
        <BookgDt><Dt>2026-10-04</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs><EndToEndId>TXN-SNB-103</EndToEndId></Refs>
            <RmtInf>
              <Ustrd>SADAD Bill Payment Code 002 Saudi Electricity Company SEC</Ustrd>
            </RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
      <!-- Entry 4: Commercial Rent Contract CONT-2026-089 -->
      <Ntry>
        <Amt Ccy="SAR">30000.00</Amt>
        <CdtDbtInd>DBIT</CdtDbtInd>
        <BookgDt><Dt>2026-10-05</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs><EndToEndId>TXN-SNB-104</EndToEndId></Refs>
            <RmtInf>
              <Ustrd>Commercial Property Rent Lease Contract CONT-2026-089</Ustrd>
            </RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
      <!-- Entry 5: Client collection -->
      <Ntry>
        <Amt Ccy="SAR">85000.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <BookgDt><Dt>2026-10-06</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Refs><EndToEndId>TXN-SNB-105</EndToEndId></Refs>
            <RmtInf>
              <Ustrd>Customer Payment Settlement #9981</Ustrd>
            </RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>
"""
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(xml_content.strip())

def generate_pdf_sample(filepath: str):
    doc = SimpleDocTemplate(filepath, pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=6
    )
    elements.append(Paragraph("<b>Riyad Bank - Corporate Account Statement</b>", title_style))
    elements.append(Paragraph("Account No: 3001122334455 | Period: October 2026 | Currency: SAR", styles['Normal']))
    elements.append(Spacer(1, 15))

    data = [
        ["Date", "Description", "Debit", "Credit", "Balance"],
        ["2026-10-07", "Outward Transfer to Al Rajhi 1002345678901", "50,000.00", "0.00", "250,000.00"],
        ["2026-10-08", "Payroll Salary Transfer Khalid Omar Al-Harbi 1003", "9,500.00", "0.00", "240,500.00"],
        ["2026-10-09", "SADAD Payment 050 GOSI Social Insurance", "7,400.00", "0.00", "233,100.00"],
        ["2026-10-10", "Rent Payment for Warehouse Contract 88992211", "18,000.00", "0.00", "215,100.00"],
        ["2026-10-11", "IT Software Cloud Subscription", "3,200.00", "0.00", "211,900.00"],
        ["2026-10-12", "Client Retainer Deposit - Riyadh Branch", "0.00", "65,000.00", "276,900.00"]
    ]

    table = Table(data, colWidths=[90, 310, 100, 100, 110])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (2, 0), (-1, -1), 'RIGHT'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))

    elements.append(table)
    doc.build(elements)

def generate_all_samples():
    os.makedirs(SAMPLE_DIR, exist_ok=True)
    excel_p = os.path.join(SAMPLE_DIR, "sample_alrajhi_statement.xlsx")
    xml_p = os.path.join(SAMPLE_DIR, "sample_snb_statement.xml")
    pdf_p = os.path.join(SAMPLE_DIR, "sample_riyad_statement.pdf")

    generate_excel_sample(excel_p)
    generate_xml_sample(xml_p)
    generate_pdf_sample(pdf_p)

    return [excel_p, xml_p, pdf_p]

if __name__ == "__main__":
    paths = generate_all_samples()
    print(f"Generated {len(paths)} sample files.")
